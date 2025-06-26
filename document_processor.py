
import os
import json
import pandas as pd
from typing import List, Dict, Any, Optional
import streamlit as st
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema import Document
import logging
from datetime import datetime

# File format specific imports
try:
    from pypdf import PdfReader
except ImportError:
    st.error("PDF processing libraries not installed. Please install pypdf.")

try:
    from docx import Document as DocxDocument
except ImportError:
    st.error("DOCX processing library not installed. Please install python-docx.")

try:
    from pptx import Presentation
except ImportError:
    st.error("PPTX processing library not installed. Please install python-pptx.")

try:
    from bs4 import BeautifulSoup
except ImportError:
    st.error("HTML processing library not installed. Please install beautifulsoup4.")

logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Process various document formats and convert to Langchain Documents"""
    
    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200):
        self.content_formats = {
            'txt': 'plain_text',
            'md': 'markdown',
            'pdf': 'pdf',
            'docx': 'word',
            'doc': 'word',
            'csv': 'tabular',
            'xls': 'tabular',
            'xlsx': 'tabular',
            'json': 'structured',
            'xml': 'structured',
            'html': 'markup',
            'htm': 'markup',
            'ppt': 'presentation',
            'pptx': 'presentation'
        }
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", " ", ""]
        )
    
    def process_file(self, file_input, file_extension: str = None, metadata: dict = None) -> List[Document]:
        """Process file based on its format
        
        Args:
            file_input: Either a file path (str) or an uploaded file object
            file_extension: Optional file extension override
            metadata: Optional metadata dictionary to merge with default metadata
        """
        try:
            # Handle both file paths and uploaded files
            if isinstance(file_input, str):
                filename = os.path.basename(file_input)
                filesize = os.path.getsize(file_input)
            else:
                filename = file_input.name
                filesize = file_input.size
                file_input.seek(0)  # Reset file pointer
            
            # Get file extension
            if not file_extension:
                file_extension = filename.split('.')[-1].lower()
            
            # Create base metadata
            base_metadata = {
                "filename": filename,
                "file_type": file_extension,
                "file_size": filesize
            }
            
            # Merge with provided metadata if any
            if metadata:
                base_metadata.update(metadata)
            
            # Map file extensions to processor methods
            extension_map = {
                'pdf': self._process_pdf,
                'doc': self._process_docx,
                'docx': self._process_docx,
                'txt': self._process_txt,
                'csv': self._process_csv,
                'xls': self._process_excel,
                'xlsx': self._process_excel,
                'ppt': self._process_pptx,
                'pptx': self._process_pptx,
                'json': self._process_json,
                'html': self._process_html,
                'htm': self._process_html,
                'xml': self._process_xml
            }
            
            # Get the appropriate processor method
            processor = extension_map.get(file_extension)
            if processor:
                text = processor(file_input)
            else:
                raise ValueError(f"Unsupported file format: {file_extension}")
            
            if not text.strip():
                raise ValueError("No text content extracted from file")
            
            # Split text into chunks
            chunks = self.text_splitter.split_text(text)
            
            # Create Document objects with enhanced metadata
            documents = []
            for i, chunk in enumerate(chunks):
                chunk_metadata = base_metadata.copy()
                chunk_metadata.update({
                    "chunk_id": i,
                    "total_chunks": len(chunks),
                    "content_type": file_extension,
                    "content_format": self._get_content_format(file_extension),
                    "chunk_size": len(chunk),
                    "is_structured": file_extension in ['csv', 'json', 'xml'],
                    "processing_timestamp": datetime.now().isoformat()
                })
                documents.append(Document(page_content=chunk, metadata=chunk_metadata))
            
            filename = file_input.name if not isinstance(file_input, str) else os.path.basename(file_input)
            logger.info(f"Processed {filename}: {len(documents)} chunks created")
            return documents
            
        except Exception as e:
            filename = file_input.name if not isinstance(file_input, str) else os.path.basename(file_input)
            logger.error(f"Error processing file {filename}: {str(e)}")
            raise e
    
    def _process_pdf(self, file_input) -> str:
        """Process PDF file"""
        try:
            # Handle both file paths and uploaded files
            if isinstance(file_input, str):
                pdf_reader = PdfReader(file_input)
            else:
                pdf_reader = PdfReader(file_input)
                file_input.seek(0)  # Reset file pointer
            
            text = ""
            for page_num, page in enumerate(pdf_reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text += f"\n--- Page {page_num + 1} ---\n{page_text}\n"
            return text
        except Exception as e:
            logger.error(f"Error processing PDF: {str(e)}")
            raise e
    
    def _process_docx(self, file_input) -> str:
        """Process DOCX file"""
        try:
            # Handle both file paths and uploaded files
            if isinstance(file_input, str):
                doc = DocxDocument(file_input)
            else:
                doc = DocxDocument(file_input)
                file_input.seek(0)  # Reset file pointer
            
            text = ""
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"
            return text
        except Exception as e:
            logger.error(f"Error processing DOCX: {str(e)}")
            raise e
    
    def _process_txt(self, file_input) -> str:
        """Process TXT file"""
        try:
            # Handle both file paths and uploaded files
            if isinstance(file_input, str):
                with open(file_input, 'r', encoding='utf-8') as f:
                    text = f.read()
            else:
                try:
                    text = file_input.read().decode('utf-8')
                except UnicodeDecodeError:
                    # Try with different encodings if UTF-8 fails
                    file_input.seek(0)  # Reset file pointer
                    text = file_input.read().decode('latin-1')
                file_input.seek(0)  # Reset file pointer
            return text
        except Exception as e:
            logger.error(f"Error processing TXT: {str(e)}")
            raise e
    
    def _process_csv(self, file_input) -> str:
        """Process CSV file"""
        try:
            # Handle both file paths and uploaded files
            if isinstance(file_input, str):
                df = pd.read_csv(file_input)
            else:
                df = pd.read_csv(file_input)
                file_input.seek(0)  # Reset file pointer
            
            filename = file_input.name if not isinstance(file_input, str) else os.path.basename(file_input)
            text = f"CSV File: {filename}\n"
            text += f"Columns: {', '.join(df.columns.tolist())}\n"
            text += f"Number of rows: {len(df)}\n\n"
            
            # Convert to string representation
            text += "Data:\n"
            text += df.to_string(index=False)
            
            return text
        except Exception as e:
            logger.error(f"Error processing CSV: {str(e)}")
            raise e
    
    def _process_excel(self, file_input) -> str:
        """Process Excel file"""
        try:
            # Handle both file paths and uploaded files
            if isinstance(file_input, str):
                excel_file = pd.ExcelFile(file_input)
            else:
                excel_file = pd.ExcelFile(file_input)
                file_input.seek(0)  # Reset file pointer
            
            filename = file_input.name if not isinstance(file_input, str) else os.path.basename(file_input)
            text = f"Excel File: {filename}\n"
            text += f"Sheets: {', '.join(excel_file.sheet_names)}\n\n"
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_input, sheet_name=sheet_name)
                text += f"\n--- Sheet: {sheet_name} ---\n"
                text += f"Columns: {', '.join(df.columns.tolist())}\n"
                text += f"Number of rows: {len(df)}\n"
                text += df.to_string(index=False)
                text += "\n"
            
            return text
        except Exception as e:
            logger.error(f"Error processing Excel: {str(e)}")
            raise e
    
    def _process_pptx(self, file_input) -> str:
        """Process PowerPoint file"""
        try:
            # Handle both file paths and uploaded files
            if isinstance(file_input, str):
                filename = os.path.basename(file_input)
                prs = Presentation(file_input)
            else:
                filename = file_input.name
                prs = Presentation(file_input)
                file_input.seek(0)  # Reset file pointer
            
            text = f"PowerPoint File: {filename}\n\n"
            
            for i, slide in enumerate(prs.slides, 1):
                text += f"\nSlide {i}:\n"
                
                # Process title
                if slide.shapes.title:
                    text += f"Title: {slide.shapes.title.text}\n"
                
                # Process text boxes and tables
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
                    elif shape.has_table:
                        text += "\nTable Content:\n"
                        for row in shape.table.rows:
                            row_text = []
                            for cell in row.cells:
                                row_text.append(cell.text)
                            text += " | ".join(row_text) + "\n"
            
            return text
        except Exception as e:
            logger.error(f"Error processing PPTX: {str(e)}")
            raise e
    
    def _process_json(self, file_input) -> str:
        """Process JSON file"""
        try:
            # Handle both file paths and uploaded files
            if isinstance(file_input, str):
                with open(file_input, 'r', encoding='utf-8') as f:
                    json_data = json.load(f)
            else:
                json_data = json.load(file_input)
                file_input.seek(0)  # Reset file pointer
            
            filename = file_input.name if not isinstance(file_input, str) else os.path.basename(file_input)
            text = f"JSON File: {filename}\n\n"
            text += json.dumps(json_data, indent=2, ensure_ascii=False)
            return text
        except Exception as e:
            logger.error(f"Error processing JSON: {str(e)}")
            raise e
    
    def _process_html(self, file_input) -> str:
        """Process HTML file"""
        try:
            # Handle both file paths and uploaded files
            if isinstance(file_input, str):
                with open(file_input, 'r', encoding='utf-8') as f:
                    html_content = f.read()
            else:
                html_content = file_input.read().decode('utf-8')
                file_input.seek(0)  # Reset file pointer
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            
            # Get text content
            text = soup.get_text()
            
            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            
            filename = file_input.name if not isinstance(file_input, str) else os.path.basename(file_input)
            return f"HTML File: {filename}\n\n{text}"
        except Exception as e:
            logger.error(f"Error processing HTML: {str(e)}")
            raise e
    
    def _process_xml(self, file_input) -> str:
        """Process XML file"""
        try:
            # Handle both file paths and uploaded files
            if isinstance(file_input, str):
                with open(file_input, 'r', encoding='utf-8') as f:
                    xml_content = f.read()
            else:
                xml_content = file_input.read().decode('utf-8')
                file_input.seek(0)  # Reset file pointer
            soup = BeautifulSoup(xml_content, 'xml')
            
            # Get text content
            text = soup.get_text()
            
            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            text = '\n'.join(line for line in lines if line)
            
            filename = file_input.name if not isinstance(file_input, str) else os.path.basename(file_input)
            return f"XML File: {filename}\n\n{text}"
        except Exception as e:
            logger.error(f"Error processing XML: {str(e)}")
            raise e
    
    def _get_content_format(self, extension: str) -> str:
        """Get standardized content format for file extension"""
        return self.content_formats.get(extension.lower(), 'unknown')

    def get_document_summary(self, documents: List[Document]) -> Dict[str, Any]:
        """Get summary statistics for processed documents"""
        if not documents:
            return {}
        
        total_chars = sum(len(doc.page_content) for doc in documents)
        total_words = sum(len(doc.page_content.split()) for doc in documents)
        
        # Get unique files and formats
        unique_files = set()
        format_counts = {}
        content_types = {}
        
        for doc in documents:
            metadata = doc.metadata
            filename = metadata.get('filename', 'Unknown')
            unique_files.add(filename)
            
            content_format = metadata.get('content_format', 'unknown')
            format_counts[content_format] = format_counts.get(content_format, 0) + 1
            
            content_type = metadata.get('content_type', 'unknown')
            content_types[content_type] = content_types.get(content_type, 0) + 1
        
        return {
            "total_documents": len(documents),
            "total_characters": total_chars,
            "total_words": total_words,
            "unique_files": len(unique_files),
            "average_chunk_size": total_chars // len(documents) if documents else 0,
            "files": list(unique_files),
            "format_distribution": format_counts,
            "content_types": content_types,
            "has_structured_content": any(doc.metadata.get('is_structured', False) for doc in documents)
        }
